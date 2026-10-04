from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Cm, Inches, Pt, RGBColor
from pptx import Presentation
from pptx.dml.color import RGBColor as PptRGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches as PptInches, Pt as PptPt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "submission"
PROJECT_MD = OUT / "01_长安食游图鉴_项目说明文档.md"
PPT_SCRIPT = OUT / "02_长安食游图鉴_路演PPT_逐页讲稿.md"

INK = PptRGBColor(5, 8, 7)
INK_2 = PptRGBColor(17, 24, 23)
PAPER = PptRGBColor(247, 242, 230)
PAPER_MUTED = PptRGBColor(190, 184, 166)
GOLD = PptRGBColor(215, 168, 85)
CINNABAR = PptRGBColor(181, 47, 53)
JADE = PptRGBColor(97, 208, 190)


def set_run_font(run, name="Microsoft YaHei", size=18, color=PAPER, bold=False):
    run.font.name = name
    run.font.size = PptPt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_box(slide, left, top, width, height, fill=INK_2, line=GOLD, radius=False):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = PptPt(0.8)
    return shape


def add_text(slide, text, left, top, width, height, size=18, color=PAPER, bold=False,
             align=PP_ALIGN.LEFT, font="Microsoft YaHei", margin=0.08):
    box = slide.shapes.add_textbox(left, top, width, height)
    box.text_frame.clear()
    box.text_frame.margin_left = PptInches(margin)
    box.text_frame.margin_right = PptInches(margin)
    box.text_frame.margin_top = PptInches(margin)
    box.text_frame.margin_bottom = PptInches(margin)
    box.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = box.text_frame.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run_font(run, font, size, color, bold)
    return box


def add_bullets(slide, items, left, top, width, height, size=18, color=PAPER_MUTED, gap=7):
    box = slide.shapes.add_textbox(left, top, width, height)
    box.text_frame.clear()
    box.text_frame.word_wrap = True
    box.text_frame.margin_left = PptInches(0.05)
    box.text_frame.margin_right = PptInches(0.05)
    box.text_frame.margin_top = PptInches(0.05)
    box.text_frame.margin_bottom = PptInches(0.05)
    for index, item in enumerate(items):
        p = box.text_frame.paragraphs[0] if index == 0 else box.text_frame.add_paragraph()
        p.text = f"· {item}"
        p.space_after = PptPt(gap)
        p.alignment = PP_ALIGN.LEFT
        for run in p.runs:
            set_run_font(run, size=size, color=color)
    return box


def add_bg(slide, number, kicker, title, subtitle=None):
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = INK
    add_text(slide, kicker.upper(), PptInches(0.65), PptInches(0.34), PptInches(4.2), PptInches(0.3), 10, CINNABAR, True, font="Arial")
    add_text(slide, title, PptInches(0.62), PptInches(0.76), PptInches(11.2), PptInches(0.66), 29, PAPER, True, font="Microsoft YaHei")
    if subtitle:
        add_text(slide, subtitle, PptInches(0.66), PptInches(1.43), PptInches(11.2), PptInches(0.38), 12, PAPER_MUTED)
    slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, PptInches(0.66), PptInches(1.93), PptInches(11.95), PptInches(0.012)).fill.solid()
    rule = slide.shapes[-1]
    rule.fill.fore_color.rgb = GOLD
    rule.line.fill.background()
    add_text(slide, f"{number:02d}  /  长安食游图鉴", PptInches(10.25), PptInches(7.08), PptInches(2.25), PptInches(0.22), 9, PAPER_MUTED, align=PP_ALIGN.RIGHT, font="Arial")


def add_chip(slide, text, left, top, width, color=GOLD):
    add_box(slide, left, top, width, PptInches(0.38), fill=INK_2, line=color, radius=True)
    add_text(slide, text, left, top + PptInches(0.01), width, PptInches(0.32), 11, color, True, align=PP_ALIGN.CENTER)


def add_flow_node(slide, label, x, y, w, accent=GOLD):
    add_box(slide, PptInches(x), PptInches(y), PptInches(w), PptInches(0.78), fill=INK_2, line=accent, radius=True)
    add_text(slide, label, PptInches(x), PptInches(y + 0.08), PptInches(w), PptInches(0.52), 17, PAPER, True, align=PP_ALIGN.CENTER)


def add_arrow(slide, x, y, w=0.45):
    add_text(slide, "→", PptInches(x), PptInches(y), PptInches(w), PptInches(0.45), 26, GOLD, True, align=PP_ALIGN.CENTER, font="Arial")


def build_ppt():
    prs = Presentation()
    prs.slide_width = PptInches(13.333)
    prs.slide_height = PptInches(7.5)
    blank = prs.slide_layouts[6]

    # 1 cover
    slide = prs.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = INK
    add_text(slide, "2026 西客松 · AI 软件赛道", PptInches(0.76), PptInches(0.62), PptInches(6), PptInches(0.3), 12, CINNABAR, True, font="Arial")
    add_text(slide, "长安食游图鉴", PptInches(0.72), PptInches(1.42), PptInches(8.8), PptInches(1.05), 44, PAPER, True)
    add_text(slide, "先看懂一味，再走进一街", PptInches(0.78), PptInches(2.55), PptInches(6.8), PptInches(0.48), 20, GOLD, False, font="Microsoft YaHei")
    add_box(slide, PptInches(0.8), PptInches(3.45), PptInches(5.2), PptInches(0.012), fill=GOLD, line=GOLD)
    add_text(slide, "西安地方美食文化图鉴 + 在地探索决策工具", PptInches(0.78), PptInches(3.72), PptInches(7.5), PptInches(0.4), 16, PAPER_MUTED)
    add_chip(slide, "菜品图鉴", PptInches(0.8), PptInches(5.32), PptInches(1.25), GOLD)
    add_chip(slide, "真实地点", PptInches(2.18), PptInches(5.32), PptInches(1.25), JADE)
    add_chip(slide, "路线决策", PptInches(3.56), PptInches(5.32), PptInches(1.25), CINNABAR)
    add_text(slide, "长安 · Hackers", PptInches(0.8), PptInches(6.65), PptInches(3), PptInches(0.3), 12, PAPER_MUTED)
    logo = ROOT / "public" / "brand" / "xihack-logo.png"
    if logo.exists():
        slide.shapes.add_picture(str(logo), PptInches(10.7), PptInches(5.6), height=PptInches(0.7))
    add_text(slide, "01", PptInches(12.25), PptInches(6.8), PptInches(0.4), PptInches(0.25), 10, PAPER_MUTED, align=PP_ALIGN.RIGHT, font="Arial")

    # 2 problem
    slide = prs.slides.add_slide(blank)
    add_bg(slide, 2, "为什么做", "游客知道想吃什么，却不知道去哪吃", "从分散信息到可执行决策，中间缺少一条可信的链路")
    cards = [
        ("看不懂", "只知道菜名，不知道历史、食材和吃法"),
        ("找不到", "美食街、老店、巷子小店和街边摊难比较"),
        ("串不起来", "时间、预算、距离和排队风险需要手动拼接"),
    ]
    for i, (head, body) in enumerate(cards):
        x = 0.78 + i * 4.1
        add_box(slide, PptInches(x), PptInches(2.55), PptInches(3.58), PptInches(2.2), fill=INK_2, line=[CINNABAR, GOLD, JADE][i], radius=True)
        add_text(slide, f"0{i + 1}", PptInches(x + 0.2), PptInches(2.78), PptInches(0.5), PptInches(0.25), 12, [CINNABAR, GOLD, JADE][i], True, font="Arial")
        add_text(slide, head, PptInches(x + 0.2), PptInches(3.2), PptInches(3.1), PptInches(0.4), 24, PAPER, True)
        add_text(slide, body, PptInches(x + 0.2), PptInches(3.78), PptInches(3.1), PptInches(0.6), 14, PAPER_MUTED)
    add_text(slide, "现有工具分别解决点评、地图或攻略，用户仍然需要自己完成最后的选择。", PptInches(0.85), PptInches(5.55), PptInches(11.4), PptInches(0.5), 19, GOLD, True, align=PP_ALIGN.CENTER)

    # 3 solution
    slide = prs.slides.add_slide(blank)
    add_bg(slide, 3, "怎么解决", "把地方文化变成一次可执行的寻味路径")
    labels = [("理解一味", CINNABAR), ("找到一处", GOLD), ("选对吃法", JADE), ("规划一程", CINNABAR)]
    for i, (label, accent) in enumerate(labels):
        x = 0.85 + i * 3.05
        add_flow_node(slide, label, x, 3.0, 2.1, accent)
        if i < len(labels) - 1:
            add_arrow(slide, x + 2.17, 3.15)
    add_text(slide, "一道菜可以关联多个地点、多种消费方式和一条路线。", PptInches(1.3), PptInches(4.65), PptInches(10.6), PptInches(0.45), 21, PAPER, True, align=PP_ALIGN.CENTER)
    add_text(slide, "用户可以只买一道菜，也可以在不同地点分散购买多种小吃。", PptInches(1.3), PptInches(5.25), PptInches(10.6), PptInches(0.4), 14, PAPER_MUTED, align=PP_ALIGN.CENTER)

    # 4 product
    slide = prs.slides.add_slide(blank)
    add_bg(slide, 4, "产品入口", "先看一味，再进入一街", "图片化菜品图鉴、历史资料、食材做法和来源集中在一个入口")
    add_box(slide, PptInches(0.8), PptInches(2.45), PptInches(5.1), PptInches(3.75), fill=INK_2, line=GOLD, radius=True)
    add_text(slide, "长安风味图鉴", PptInches(1.08), PptInches(2.8), PptInches(3.2), PptInches(0.4), 22, GOLD, True)
    add_text(slide, "肉夹馍   羊肉泡馍   凉皮   油泼面", PptInches(1.08), PptInches(3.55), PptInches(4.35), PptInches(0.4), 17, PAPER, True)
    add_text(slide, "甑糕   葫芦鸡   臊子面   水盆羊肉", PptInches(1.08), PptInches(4.08), PptInches(4.35), PptInches(0.4), 15, PAPER_MUTED)
    add_text(slide, "历史 · 食材 · 调味 · 做法 · 公开来源", PptInches(1.08), PptInches(4.95), PptInches(4.35), PptInches(0.35), 13, JADE)
    add_chip(slide, "14 道菜品资料", PptInches(1.08), PptInches(5.48), PptInches(1.65), CINNABAR)
    add_box(slide, PptInches(6.45), PptInches(2.45), PptInches(5.9), PptInches(3.75), fill=INK_2, line=CINNABAR, radius=True)
    add_text(slide, "一味小史", PptInches(6.8), PptInches(2.82), PptInches(2.1), PptInches(0.35), 20, GOLD, True)
    add_bullets(slide, ["菜品历史与百度百科公开资料链接", "食材、调味和制作步骤结构化展示", "图片优先复用更新版网页图片逻辑", "缺图回退公开媒体库，不生成虚构事实"], PptInches(6.8), PptInches(3.38), PptInches(4.9), PptInches(2.0), 15)

    # 5 place
    slide = prs.slides.add_slide(blank)
    add_bg(slide, 5, "从菜到店", "把“想吃”变成“知道去哪”")
    add_bullets(slide, ["店铺名称、区域、地址和地图点位", "人均价格、菜系、营业时间和服务方式", "堂食、外带、交通指南等消费标签", "菜单与菜品图鉴双向关联", "开启定位后展示用户与门店距离"], PptInches(0.9), PptInches(2.55), PptInches(5.3), PptInches(3.1), 17, PAPER)
    add_box(slide, PptInches(7.0), PptInches(2.45), PptInches(4.85), PptInches(3.8), fill=INK_2, line=JADE, radius=True)
    add_text(slide, "沿街拾味", PptInches(7.35), PptInches(2.82), PptInches(2.2), PptInches(0.35), 22, JADE, True)
    add_text(slide, "秦缘肉夹馍（骡马市总店）", PptInches(7.35), PptInches(3.55), PptInches(4.0), PptInches(0.35), 15, PAPER, True)
    add_text(slide, "碑林区 · 骡马市   人均 ¥17", PptInches(7.35), PptInches(4.02), PptInches(3.8), PptInches(0.3), 13, PAPER_MUTED)
    add_text(slide, "堂食 / 外带 / 交通指南", PptInches(7.35), PptInches(4.52), PptInches(3.8), PptInches(0.3), 13, GOLD)
    add_text(slide, "定位后显示距离 · 资料来自游陕西截图", PptInches(7.35), PptInches(5.18), PptInches(4.0), PptInches(0.35), 11, PAPER_MUTED)

    # 6 route
    slide = prs.slides.add_slide(blank)
    add_bg(slide, 6, "路线 Demo", "一句话生成一条可走的美食路线", "Agent 理解条件，规则层负责可行性")
    add_box(slide, PptInches(0.82), PptInches(2.43), PptInches(4.3), PptInches(3.8), fill=INK_2, line=GOLD, radius=True)
    add_text(slide, "钟楼附近 · 2 小时 · 3 种小吃", PptInches(1.15), PptInches(2.94), PptInches(3.65), PptInches(0.5), 19, PAPER, True)
    add_text(slide, "不想去连锁店，优先巷子和老店", PptInches(1.15), PptInches(3.62), PptInches(3.4), PptInches(0.35), 14, PAPER_MUTED)
    add_chip(slide, "落笔成行", PptInches(1.15), PptInches(4.45), PptInches(1.35), CINNABAR)
    add_text(slide, "地点 · 时长 · 预算 · 口味 · 消费方式", PptInches(1.15), PptInches(5.3), PptInches(3.55), PptInches(0.38), 12, GOLD)
    add_flow_node(slide, "解析条件", 5.8, 2.72, 1.6, CINNABAR)
    add_arrow(slide, 7.5, 2.87, 0.44)
    add_flow_node(slide, "查菜与店", 8.0, 2.72, 1.6, GOLD)
    add_arrow(slide, 9.7, 2.87, 0.44)
    add_flow_node(slide, "算预算距离", 10.2, 2.72, 1.6, JADE)
    add_text(slide, "路线结果：节点顺序 · 预计时长 · 预算 · 风险 · Plan B", PptInches(5.82), PptInches(4.25), PptInches(6.0), PptInches(0.45), 16, PAPER, True, align=PP_ALIGN.CENTER)
    add_text(slide, "模型失败、地图失败或网络失败时，回退到本地资料。", PptInches(5.95), PptInches(5.05), PptInches(5.8), PptInches(0.4), 13, PAPER_MUTED, align=PP_ALIGN.CENTER)

    # 7 trust
    slide = prs.slides.add_slide(blank)
    add_bg(slide, 7, "可信数据", "来源和时间标记，是产品的一部分")
    rows = [
        ("游陕西小程序", "菜品、店铺、榜单截图", CINNABAR),
        ("百度百科词条", "历史与基本形制参考", GOLD),
        ("更新版网页", "菜品图片提示与展示参考", JADE),
        ("Wikimedia Commons", "缺图时的公开媒体库候选", PAPER),
    ]
    for i, (source, desc, accent) in enumerate(rows):
        y = 2.38 + i * 0.73
        add_box(slide, PptInches(0.9), PptInches(y), PptInches(3.0), PptInches(0.5), fill=INK_2, line=accent, radius=True)
        add_text(slide, source, PptInches(1.08), PptInches(y + 0.05), PptInches(2.6), PptInches(0.35), 14, accent, True)
        add_text(slide, desc, PptInches(4.25), PptInches(y + 0.05), PptInches(5.5), PptInches(0.35), 14, PAPER_MUTED)
    add_box(slide, PptInches(9.55), PptInches(2.45), PptInches(2.5), PptInches(2.85), fill=INK_2, line=GOLD, radius=True)
    add_text(slide, "数据状态", PptInches(9.9), PptInches(2.84), PptInches(1.8), PptInches(0.32), 20, GOLD, True, align=PP_ALIGN.CENTER)
    add_text(slide, "已核验\n待确认\n过期 / 需更新\n演示 seed", PptInches(10.0), PptInches(3.5), PptInches(1.6), PptInches(1.4), 15, PAPER, True, align=PP_ALIGN.CENTER)
    add_text(slide, "没有来源的信息不自动补全。", PptInches(1.0), PptInches(5.55), PptInches(10.6), PptInches(0.4), 19, CINNABAR, True, align=PP_ALIGN.CENTER)

    # 8 architecture
    slide = prs.slides.add_slide(blank)
    add_bg(slide, 8, "技术实现", "AI 负责理解，程序负责校验")
    add_flow_node(slide, "自然语言输入", 0.95, 2.55, 2.0, CINNABAR)
    add_arrow(slide, 3.08, 2.7)
    add_flow_node(slide, "Agent 解析", 3.6, 2.55, 1.8, GOLD)
    add_arrow(slide, 5.54, 2.7)
    add_flow_node(slide, "结构化数据", 6.05, 2.55, 1.9, JADE)
    add_arrow(slide, 8.1, 2.7)
    add_flow_node(slide, "规则校验", 8.62, 2.55, 1.8, GOLD)
    add_arrow(slide, 10.56, 2.7)
    add_flow_node(slide, "路线结果", 11.05, 2.55, 1.55, CINNABAR)
    add_bullets(slide, ["React 18 + Vite：可运行前端原型", "本地 seed + source registry：断网可演示", "routeApi：预算、距离和时间计算", "agentApi：自然语言解析和规则回退", "TianDiMap：地图可用时展示门店点位"], PptInches(1.05), PptInches(4.2), PptInches(5.2), PptInches(1.85), 15)
    add_bullets(slide, ["价格不由模型猜测", "营业状态保留时间标记", "网络失败有降级路径", "第三方资源记录来源和许可证"], PptInches(7.0), PptInches(4.2), PptInches(4.9), PptInches(1.85), 15, PAPER_MUTED)

    # 9 value
    slide = prs.slides.add_slide(blank)
    add_bg(slide, 9, "价值与创新", "把地方文化内容转化为真实消费决策")
    add_box(slide, PptInches(0.85), PptInches(2.45), PptInches(5.55), PptInches(3.7), fill=INK_2, line=CINNABAR, radius=True)
    add_text(slide, "差异化", PptInches(1.2), PptInches(2.85), PptInches(1.4), PptInches(0.32), 21, CINNABAR, True)
    add_bullets(slide, ["菜品文化图鉴连接真实地点", "一道菜支持多种购买方式", "多地点路线与 Plan B", "来源、时间和待确认状态可见"], PptInches(1.2), PptInches(3.45), PptInches(4.65), PptInches(1.9), 16)
    add_box(slide, PptInches(6.95), PptInches(2.45), PptInches(5.4), PptInches(3.7), fill=INK_2, line=JADE, radius=True)
    add_text(slide, "潜在客户", PptInches(7.3), PptInches(2.85), PptInches(1.6), PptInches(0.32), 21, JADE, True)
    add_bullets(slide, ["酒店与民宿：本地接待方案", "旅行社与地接：美食游路线", "餐厅：特色菜展示与客流", "美食街：文化传播与消费组织"], PptInches(7.3), PptInches(3.45), PptInches(4.5), PptInches(1.9), 16)

    # 10 close
    slide = prs.slides.add_slide(blank)
    add_bg(slide, 10, "落地计划", "让每一味长安风味，都有来处也有去处")
    add_flow_node(slide, "P0\n真实来源 + 现场 Demo", 1.1, 2.75, 2.7, CINNABAR)
    add_flow_node(slide, "P1\n商家维护 + 多路线比较", 5.25, 2.75, 2.7, GOLD)
    add_flow_node(slide, "P2\n审核后台 + 用户投稿", 9.4, 2.75, 2.7, JADE)
    add_text(slide, "今天的 Demo 已经跑通：看菜 → 看店 → 定位 → 规划路线。", PptInches(1.1), PptInches(4.55), PptInches(11.1), PptInches(0.45), 22, PAPER, True, align=PP_ALIGN.CENTER)
    add_text(slide, "长安食游图鉴 · 先看懂一味，再走进一街", PptInches(1.1), PptInches(5.35), PptInches(11.1), PptInches(0.4), 17, GOLD, False, align=PP_ALIGN.CENTER)

    output = OUT / "02_长安食游图鉴_路演PPT.pptx"
    prs.save(output)


def add_docx_paragraph(doc, text, style=None, bold=False):
    p = doc.add_paragraph(style=style)
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(10.5)
    return p


def build_docx():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.1)
    section.right_margin = Cm(2.1)
    styles = doc.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"].font.size = Pt(10.5)
    styles["Title"].font.name = "Microsoft YaHei"
    styles["Title"].font.size = Pt(24)
    styles["Heading 1"].font.name = "Microsoft YaHei"
    styles["Heading 1"].font.size = Pt(16)
    styles["Heading 1"].font.color.rgb = RGBColor(181, 47, 53)
    styles["Heading 2"].font.name = "Microsoft YaHei"
    styles["Heading 2"].font.size = Pt(13)
    styles["Heading 2"].font.color.rgb = RGBColor(145, 103, 48)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("长安食游图鉴")
    run.bold = True
    run.font.color.rgb = RGBColor(181, 47, 53)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("项目说明文档｜2026 西客松 AI 软件赛道").italic = True

    lines = PROJECT_MD.read_text(encoding="utf-8").splitlines()
    in_code = False
    table_rows = []

    def flush_table():
        nonlocal table_rows
        if not table_rows:
            return
        cols = max(len(row) for row in table_rows)
        table = doc.add_table(rows=1, cols=cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = "Table Grid"
        for j, cell in enumerate(table.rows[0].cells):
            cell.text = table_rows[0][j] if j < len(table_rows[0]) else ""
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for row in table_rows[1:]:
            cells = table.add_row().cells
            for j, cell in enumerate(cells):
                cell.text = row[j] if j < len(row) else ""
        table_rows = []

    for line in lines:
        if line.startswith("# 长安食游图鉴"):
            continue
        if line.startswith("```"):
            flush_table()
            in_code = not in_code
            continue
        if in_code:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            r = p.add_run(line)
            r.font.name = "Consolas"
            r.font.size = Pt(9)
            continue
        if line.startswith("## "):
            flush_table()
            doc.add_heading(line[3:], level=1)
        elif line.startswith("### "):
            flush_table()
            doc.add_heading(line[4:], level=2)
        elif line.startswith("- "):
            flush_table()
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(line[2:])
        elif line[:2].isdigit() and line[2:4] == ". ":
            flush_table()
            p = doc.add_paragraph(style="List Number")
            p.add_run(line[4:])
        elif line.startswith("|"):
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            if all(set(c) <= set("-: ") for c in cells):
                continue
            table_rows.append(cells)
        elif line.strip() == "":
            flush_table()
            continue
        else:
            flush_table()
            add_docx_paragraph(doc, line)
    flush_table()
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("长安食游图鉴 · 项目说明文档 · 2026 西客松").font.size = Pt(8)
    doc.save(OUT / "01_长安食游图鉴_项目说明文档.docx")


if __name__ == "__main__":
    build_ppt()
    build_docx()
    print("generated", OUT / "02_长安食游图鉴_路演PPT.pptx")
    print("generated", OUT / "01_长安食游图鉴_项目说明文档.docx")
